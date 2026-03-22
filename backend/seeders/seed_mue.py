"""
Seed Snowflake mue_edits from CMS NCCI MUE Q2 2026 files.

Direct downloads — no license gate. URLs confirmed from live CMS page.
Files contain Excel (.xlsx) spreadsheets inside each ZIP.

Effective April 1, 2026.
"""
import csv
import io
import sys
import zipfile
from pathlib import Path

import httpx
import openpyxl

sys.path.insert(0, str(Path(__file__).parent.parent))
from services.snowflake_client import get_connection

MUE_SOURCES = [
    (
        "https://www.cms.gov/files/zip/medicare-ncci-2026-q2-practitioner-services-mue-table.zip",
        "practitioner",
    ),
    (
        "https://www.cms.gov/files/zip/medicare-ncci-2026-q2-facility-outpatient-hospital-services-mue-table.zip",
        "outpatient",
    ),
    (
        "https://www.cms.gov/files/zip/medicare-ncci-2026-q2-dme-supplier-services-mue-table.zip",
        "dme",
    ),
]

BATCH = 2_000


def _parse_xlsx(content: bytes, stype: str) -> list[tuple]:
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active

    all_rows = list(ws.iter_rows(values_only=True))

    # CMS xlsx files have variable numbers of title/disclaimer rows before data.
    # Find the header row — the one containing "HCPCS" or "CPT" in any cell.
    data_start = 1  # default: skip just row 0
    for i, row in enumerate(all_rows):
        row_text = " ".join(str(c).upper() for c in row if c)
        if "HCPCS" in row_text or "CPT" in row_text:
            data_start = i + 1  # data starts after this header row
            break

    # Also detect which columns hold code / MUE value / MAI
    # by inspecting the header row
    code_col, mue_col, mai_col = 0, 1, 2
    if data_start > 0:
        header = [str(c).upper() if c else "" for c in all_rows[data_start - 1]]
        for j, h in enumerate(header):
            if "HCPCS" in h or "CPT" in h:
                code_col = j
            elif "MUE" in h and "VALUE" in h:
                mue_col = j
            elif "MAI" in h or "ADJUDICATION" in h:
                mai_col = j

    rows = []
    for row in all_rows[data_start:]:
        if not row or len(row) <= max(code_col, mue_col, mai_col):
            continue
        cpt = str(row[code_col]).strip() if row[code_col] else ""
        if not cpt or cpt.upper() in ("NONE", "NULL", ""):
            continue
        if not (cpt[0].isdigit() or cpt[0].isalpha()):
            continue
        try:
            mue_val = int(float(str(row[mue_col]))) if row[mue_col] is not None else None
            adj_ind = int(float(str(row[mai_col]))) if row[mai_col] is not None else None
        except (ValueError, TypeError):
            continue
        if mue_val is None or adj_ind is None:
            continue
        rows.append((cpt, mue_val, adj_ind, stype))
    return rows


def _parse_csv(content: bytes, stype: str) -> list[tuple]:
    rows = []
    reader = csv.reader(io.StringIO(content.decode("utf-8-sig")))
    for i, line in enumerate(reader):
        if i == 0 or not line or not line[0].strip():
            continue
        try:
            cpt = line[0].strip()
            mue_val = int(line[1].strip())
            adj_ind = int(line[2].strip())
            rows.append((cpt, mue_val, adj_ind, stype))
        except (IndexError, ValueError):
            continue
    return rows


def _batch_insert(cur, rows: list[tuple]) -> int:
    total = 0
    for i in range(0, len(rows), BATCH):
        chunk = rows[i: i + BATCH]
        placeholders = ",".join(["(%s,%s,%s,%s)"] * len(chunk))
        flat = [v for row in chunk for v in row]
        cur.execute(
            "INSERT INTO mue_edits "
            "(cpt_code, mue_value, adjudication_indicator, service_type) "
            f"VALUES {placeholders}",
            flat,
        )
        total += len(chunk)
    return total


def seed():
    conn = get_connection()
    cur = conn.cursor()
    cur.execute("USE DATABASE vinci_db")
    cur.execute("USE SCHEMA insurer_policies")
    cur.execute("DELETE FROM mue_edits")

    total = 0
    for url, stype in MUE_SOURCES:
        print(f"  Downloading MUE {stype} …")
        try:
            resp = httpx.get(url, follow_redirects=True, timeout=120)
            resp.raise_for_status()
        except Exception as e:
            print(f"  WARNING: {e}")
            continue

        with zipfile.ZipFile(io.BytesIO(resp.content)) as zf:
            xl_files = [n for n in zf.namelist() if n.lower().endswith((".xlsx", ".xls"))]
            csv_files = [n for n in zf.namelist() if n.lower().endswith(".csv")]
            targets = xl_files or csv_files

            for fname in targets:
                content = zf.read(fname)
                if fname.lower().endswith((".xlsx", ".xls")):
                    rows = _parse_xlsx(content, stype)
                else:
                    rows = _parse_csv(content, stype)
                n = _batch_insert(cur, rows)
                total += n
                print(f"    {fname}: {n:,} rows")

    conn.commit()
    conn.close()
    print(f"  Total MUE edits seeded: {total:,} ✓")


if __name__ == "__main__":
    seed()
